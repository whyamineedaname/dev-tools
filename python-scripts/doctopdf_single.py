"""
单文件转 PDF - 供 Electron 调用
用法: python doctopdf_single.py <输入文件路径> <输出文件路径>
支持: .ppt, .pptx, .jpg, .jpeg, .png
"""
import sys
import os
import subprocess
import hashlib

# 强制 UTF-8 输出：Windows 下 Python 管道默认用 GBK/cp936，会导致 Electron 侧中文乱码
try:
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
except Exception:
    pass

# 彻底替换reportlab的md5函数
def safe_md5(data):
    try:
        return hashlib.md5(data, usedforsecurity=False).digest()
    except TypeError:
        return hashlib.md5(data).digest()

import reportlab.lib.utils
reportlab.lib.utils.md5 = safe_md5

from pathlib import Path
from pptx import Presentation
from PIL import Image
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
import warnings
warnings.filterwarnings('ignore')


class SafeImageReader(ImageReader):
    def __init__(self, image, width=None, height=None):
        if hasattr(image, 'save'):
            self.temp_path = f"temp_safe_img_{os.urandom(6).hex()}.jpg"
            image.save(self.temp_path, 'JPEG', quality=95)
            super().__init__(self.temp_path, width, height)
        else:
            super().__init__(image, width, height)

    def __del__(self):
        if hasattr(self, 'temp_path') and os.path.exists(self.temp_path):
            try:
                os.remove(self.temp_path)
            except:
                pass


def ppt_to_pdf_powerpoint(ppt_path: str, pdf_path: str) -> bool:
    """PPT/PPTX 转 PDF（PowerPoint COM）"""
    if os.name != 'nt':
        return False
    import win32com.client
    powerpoint = None
    presentation = None
    try:
        powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
        powerpoint.Visible = True
        powerpoint.DisplayAlerts = 0
        presentation = powerpoint.Presentations.Open(ppt_path)
        presentation.SaveAs(pdf_path, FileFormat=32)
        return True
    except Exception as e:
        print(f"PowerPoint 转换失败: {e}", file=sys.stderr)
        return False
    finally:
        if presentation:
            try:
                presentation.Close()
            except:
                pass
        if powerpoint:
            try:
                powerpoint.Quit()
            except:
                pass
        try:
            os.system("taskkill /f /im POWERPNT.EXE 2>nul")
        except:
            pass
        del presentation
        del powerpoint


WPS_CLI = "/Applications/wpsoffice.app/Contents/MacOS/wpscli"

SOFFICE_CANDIDATES = [
    "/opt/homebrew/bin/soffice",
    "/usr/local/bin/soffice",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
]


def find_soffice() -> str:
    """定位 LibreOffice soffice 可执行文件"""
    import shutil
    for c in SOFFICE_CANDIDATES:
        if os.path.exists(c):
            return c
    w = shutil.which("soffice")
    return w or ""


def ppt_to_pdf_libreoffice(ppt_path: str, pdf_path: str) -> bool:
    """PPT/PPTX 转 PDF（LibreOffice headless，macOS 首选，无沙盒/登录限制）"""
    soffice = find_soffice()
    if not soffice:
        return False
    try:
        out_dir = os.path.dirname(os.path.abspath(pdf_path))
        os.makedirs(out_dir, exist_ok=True)
        result = subprocess.run(
            [soffice, "--headless", "--convert-to", "pdf", "--outdir", out_dir, ppt_path],
            capture_output=True, text=True, timeout=300,
        )
        if result.returncode != 0:
            print(f"LibreOffice 转换失败(code={result.returncode}): {(result.stderr or '')[:200]}", file=sys.stderr)
            return False
        # soffice 输出文件名固定为 <源文件名>.pdf
        generated = os.path.join(out_dir, os.path.splitext(os.path.basename(ppt_path))[0] + ".pdf")
        if not os.path.exists(generated):
            print("LibreOffice 转换后未找到输出 PDF", file=sys.stderr)
            return False
        if os.path.abspath(generated) != os.path.abspath(pdf_path):
            if os.path.exists(pdf_path):
                os.remove(pdf_path)
            os.rename(generated, pdf_path)
        return True
    except Exception as e:
        print(f"LibreOffice 转换异常: {e}", file=sys.stderr)
        return False


def ppt_to_pdf_wps(ppt_path: str, pdf_path: str) -> bool:
    """PPT/PPTX 转 PDF（macOS WPS CLI）"""
    try:
        if not os.path.exists(WPS_CLI):
            return False
        # 目标文件若存在，wpscli 会自动重命名而非覆盖，先删除以保证输出路径一致
        if os.path.exists(pdf_path):
            os.remove(pdf_path)
        result = subprocess.run(
            [WPS_CLI, "ppt2pdf", ppt_path, "-o", pdf_path, "--json"],
            capture_output=True, text=True, timeout=300,
        )
        if result.returncode == 0 and os.path.exists(pdf_path):
            return True
        # 100/101 = 未登录 WPS 账号 / 无 VIP 权限
        err = (result.stderr or result.stdout or "").strip()
        if result.returncode in (100, 101):
            print("WPS 转换需要登录账号并开通 VIP 权限", file=sys.stderr)
        else:
            print(f"WPS 转换失败(code={result.returncode}): {err[:200]}", file=sys.stderr)
        return False
    except Exception as e:
        print(f"WPS 转换异常: {e}", file=sys.stderr)
        return False


def pptx_to_pdf_pure_python(pptx_path: str, pdf_path: str) -> bool:
    """PPTX 纯 Python 转换"""
    try:
        prs = Presentation(pptx_path)
        c = canvas.Canvas(pdf_path, pagesize=A4)
        a4_width, a4_height = A4
        slide_count = len(prs.slides)

        if slide_count == 0:
            return False

        for i, slide in enumerate(prs.slides):
            try:
                if hasattr(slide.shapes, '_spTree') and hasattr(slide.shapes._spTree, 'xml'):
                    slide.shapes._spTree.xml = slide.shapes._spTree.xml.replace('w:txBox', 'a:txBox')
            except:
                pass

            slide_width = prs.slide_width.inches * 72
            slide_height = prs.slide_height.inches * 72

            scale_x = a4_width / slide_width
            scale_y = a4_height / slide_height
            scale = min(scale_x, scale_y)

            x_pos = (a4_width - slide_width * scale) / 2
            y_pos = (a4_height - slide_height * scale) / 2

            c.setStrokeColorRGB(0.3, 0.3, 0.3)
            c.rect(x_pos, y_pos, slide_width * scale, slide_height * scale, fill=0)

            c.setFillColorRGB(0, 0, 0)
            c.setFont("Helvetica", 11)
            slide_title = f"幻灯片 {i+1}/{slide_count}"
            try:
                if slide.shapes.title and hasattr(slide.shapes.title, 'text'):
                    slide_title = slide.shapes.title.text[:30] + "..." if len(slide.shapes.title.text) > 30 else slide.shapes.title.text
            except:
                pass
            c.drawString(x_pos + 15, y_pos + slide_height * scale - 25, slide_title)

            c.setFont("Helvetica", 9)
            c.setFillColorRGB(0.5, 0.5, 0.5)
            file_name = os.path.basename(pptx_path)[:20] + "..." if len(os.path.basename(pptx_path)) > 20 else os.path.basename(pptx_path)
            c.drawString(x_pos + 15, y_pos + slide_height * scale - 45, f"来源：{file_name}")

            c.setFillColorRGB(0.7, 0.7, 0.7)
            c.drawString(x_pos + 15, y_pos + 20, "纯Python转换（基础格式）- 复杂内容建议安装PowerPoint")

            c.showPage()

        c.save()
        return True
    except Exception as e:
        print(f"纯Python转换异常: {e}", file=sys.stderr)
        return False


def image_to_pdf(image_path: str, pdf_path: str) -> bool:
    """图片转 PDF"""
    try:
        with Image.open(image_path) as img:
            if img.mode in ('RGBA', 'P', 'LA'):
                bg = Image.new('RGB', img.size, (255, 255, 255))
                if img.mode == 'RGBA':
                    bg.paste(img, mask=img.split()[3])
                else:
                    bg.paste(img)
                img = bg

            c = canvas.Canvas(pdf_path, pagesize=A4)
            a4_w, a4_h = A4
            img_w, img_h = img.size

            scale = min(a4_w / img_w, a4_h / img_h)
            x = (a4_w - img_w * scale) / 2
            y = (a4_h - img_h * scale) / 2

            safe_reader = SafeImageReader(img)
            c.drawImage(safe_reader, x, y, width=img_w*scale, height=img_h*scale)
            c.save()

        return True
    except Exception as e:
        print(f"图片转换异常: {e}", file=sys.stderr)
        return False


def check_powerpoint_availability() -> bool:
    """检测本机可用的 Office 应用（Windows: PowerPoint COM；macOS: WPS CLI）"""
    if os.name == 'nt':
        try:
            import win32com.client
            powerpoint = win32com.client.DispatchEx("PowerPoint.Application")
            powerpoint.Quit()
            del powerpoint
            return True
        except Exception:
            return False
    else:
        # macOS：检查 WPS 命令行工具
        return os.path.exists(WPS_CLI)


def convert_to_pdf(input_path: str, output_path: str) -> bool:
    """单文件转 PDF"""
    if not os.path.exists(input_path):
        print(f"错误: 文件不存在 - {input_path}", file=sys.stderr)
        return False

    ext = os.path.splitext(input_path)[1].lower()
    supported = ('.ppt', '.pptx', '.jpg', '.jpeg', '.png')

    if ext not in supported:
        print(f"错误: 不支持的文件格式 - {ext}", file=sys.stderr)
        return False

    # 确保输出目录存在
    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)

    if ext == '.ppt':
        if os.name == 'nt':
            return ppt_to_pdf_powerpoint(input_path, output_path)
        # macOS：.ppt 优先 LibreOffice，其次 WPS
        if ppt_to_pdf_libreoffice(input_path, output_path):
            return True
        if ppt_to_pdf_wps(input_path, output_path):
            return True
        print("LibreOffice/WPS 转换均失败，macOS 上 .ppt 暂无其他转换方案", file=sys.stderr)
        return False

    elif ext == '.pptx':
        if os.name == 'nt':
            if ppt_to_pdf_powerpoint(input_path, output_path):
                return True
            print("本机 Office 转换失败，尝试纯 Python 转换...", file=sys.stderr)
            return pptx_to_pdf_pure_python(input_path, output_path)
        # macOS：LibreOffice → WPS → 纯 Python 逐级降级
        if ppt_to_pdf_libreoffice(input_path, output_path):
            return True
        print("LibreOffice 转换失败，尝试 WPS...", file=sys.stderr)
        if ppt_to_pdf_wps(input_path, output_path):
            return True
        print("本机 LibreOffice/WPS 转换失败，尝试纯 Python 转换...", file=sys.stderr)
        return pptx_to_pdf_pure_python(input_path, output_path)

    elif ext in ('.jpg', '.jpeg', '.png'):
        return image_to_pdf(input_path, output_path)

    return False


def main():
    if len(sys.argv) < 3:
        print("用法: python doctopdf_single.py <输入文件> <输出文件>", file=sys.stderr)
        print("支持格式: .ppt, .pptx, .jpg, .jpeg, .png", file=sys.stderr)
        sys.exit(1)

    input_path = sys.argv[1]
    output_path = sys.argv[2]

    success = convert_to_pdf(input_path, output_path)
    if success:
        print(f"SUCCESS:{output_path}")
        sys.exit(0)
    else:
        print("FAILED", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main()
